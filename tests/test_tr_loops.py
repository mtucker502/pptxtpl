"""End-to-end tests for table row loops via the {%tr%} prefix."""

import os

import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from pptxtpl import PptxTemplate
from pptxtpl.exceptions import InvalidTemplateError, TemplateRenderError

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _row_texts(row):
    return [t.text for t in row.findall(f".//{A}t")]


def test_tr_for_separate_row_pattern(tmp_dir):
    """The supported pattern: for-tag and endfor-tag in separate rows.

    4-row template: header, for-row, body, endfor-row.
    The for-row and endfor-row are removed; the body row is repeated.
    """
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    table = s.shapes.add_table(
        4, 3, Inches(1), Inches(1), Inches(7), Inches(3)
    ).table
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(0, 2).text = "Status"
    table.cell(1, 0).text = "{%tr for m in metrics %}"
    table.cell(2, 0).text = "{{ m.name }}"
    table.cell(2, 1).text = "{{ m.value }}"
    table.cell(2, 2).text = "{{ m.status }}"
    table.cell(3, 0).text = "{%tr endfor %}"
    src = os.path.join(tmp_dir, "tr.pptx")
    prs.save(src)

    tpl = PptxTemplate(src)
    tpl.render(
        {
            "metrics": [
                {"name": "Revenue", "value": "$4.2M", "status": "OK"},
                {"name": "NPS", "value": "72", "status": "Good"},
                {"name": "Churn", "value": "3.1%", "status": "Risk"},
            ]
        }
    )
    out = os.path.join(tmp_dir, "tr_out.pptx")
    tpl.save(out)

    out_prs = Presentation(out)
    root = etree.fromstring(etree.tostring(out_prs.slides[0]._element))
    rows = root.findall(f".//{A}tbl/{A}tr")
    assert len(rows) == 4  # header + 3 metrics
    assert _row_texts(rows[0]) == ["Metric", "Value", "Status"]
    assert _row_texts(rows[1]) == ["Revenue", "$4.2M", "OK"]
    assert _row_texts(rows[2]) == ["NPS", "72", "Good"]
    assert _row_texts(rows[3]) == ["Churn", "3.1%", "Risk"]


def test_tr_for_same_row_raises_clear_error(tmp_dir):
    """The unsupported (same-row) pattern must raise a clear error.

    Previously this produced a confusing 'Unexpected end of template' from
    Jinja because elevation consumed the endfor along with the for.
    """
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    table = s.shapes.add_table(
        2, 3, Inches(1), Inches(1), Inches(7), Inches(2)
    ).table
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(0, 2).text = "Status"
    table.cell(1, 0).text = "{%tr for m in metrics %}{{ m.name }}"
    table.cell(1, 1).text = "{{ m.value }}"
    table.cell(1, 2).text = "{{ m.status }}{%tr endfor %}"
    src = os.path.join(tmp_dir, "tr_bad.pptx")
    prs.save(src)

    tpl = PptxTemplate(src)
    with pytest.raises((InvalidTemplateError, TemplateRenderError)) as exc:
        tpl.render({"metrics": [{"name": "n", "value": "v", "status": "s"}]})
    msg = str(exc.value)
    assert "tr" in msg and "separate" in msg.lower()
