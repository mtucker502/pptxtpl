"""Tests for precomputed autofit (shrink text on overflow)."""

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from pptxtpl.autofit import fit_presentation, fit_shape, fit_slide

pytest.importorskip("PIL", reason="autofit requires Pillow")


def _body_pr(shape):
    return shape.text_frame._txBody.bodyPr


def _norm_autofit(shape):
    return _body_pr(shape).find(qn("a:normAutofit"))


def _textbox(slide, text, *, width_in, height_in, size_pt=18):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(width_in), Inches(height_in))
    tf = box.text_frame
    tf.word_wrap = True
    tf.text = text
    for paragraph in tf.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(size_pt)
    return box


def test_overflowing_shape_gets_font_scale():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = _textbox(slide, "word " * 300, width_in=3, height_in=1)

    scale = fit_shape(box)

    assert scale is not None
    assert scale < 1.0
    node = _norm_autofit(box)
    assert node is not None
    assert int(node.get("fontScale")) < 100000


def test_short_text_is_not_shrunk():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = _textbox(slide, "Hi", width_in=8, height_in=3)

    scale = fit_shape(box)

    assert scale == 1.0
    assert _norm_autofit(box) is None


def test_single_line_taller_than_its_box_is_not_shrunk():
    # Designers routinely put 30pt text in a box barely 32pt tall. The line box
    # exceeds the frame but the glyphs do not, so it must be left alone.
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = _textbox(slide, "Customer Presentation", width_in=12, height_in=0.45, size_pt=30)

    scale = fit_shape(box)

    assert scale == 1.0
    assert _norm_autofit(box) is None


def test_wrapping_text_in_a_short_box_is_shrunk():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = _textbox(slide, "Customer Presentation " * 6, width_in=4, height_in=0.45, size_pt=30)

    scale = fit_shape(box)

    assert scale < 1.0


def test_hard_line_breaks_are_counted():
    # An <a:br/> is not a run. Measuring only run text loses every hard break
    # in a paragraph and badly undercounts bulleted blocks.
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = _textbox(slide, "Short line", width_in=8, height_in=1.5, size_pt=18)
    paragraph = box.text_frame.paragraphs[0]
    for _ in range(20):
        paragraph.add_line_break()
        run = paragraph.add_run()
        run.text = "Short line"
        run.font.size = Pt(18)

    assert fit_shape(box) < 1.0


def test_no_autofit_is_replaced():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = _textbox(slide, "word " * 300, width_in=3, height_in=1)

    body_pr = _body_pr(box)
    body_pr.append(body_pr.makeelement(qn("a:noAutofit"), {}))

    fit_shape(box)

    assert body_pr.find(qn("a:noAutofit")) is None
    assert body_pr.find(qn("a:normAutofit")) is not None


def test_scale_never_goes_below_min():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = _textbox(slide, "word " * 5000, width_in=2, height_in=0.4)

    scale = fit_shape(box, min_scale=0.5)

    assert scale == pytest.approx(0.5)


def test_empty_and_sizeless_shapes_are_skipped():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    empty = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    picture_like = slide.shapes.add_shape(1, Inches(1), Inches(3), Inches(1), Inches(1))

    assert fit_shape(empty) is None
    assert fit_shape(picture_like) is None


def test_footer_placeholders_are_skipped():
    # python-pptx does not copy footer/slide-number placeholders onto a new
    # slide, so clone one from the layout to have something to skip.
    import copy

    prs = Presentation()
    layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    footer_src = next(p for p in layout.placeholders if p.placeholder_format.idx == 11)
    slide.shapes._spTree.append(copy.deepcopy(footer_src._element))
    footer = next(p for p in slide.placeholders if p.placeholder_format.idx == 11)
    footer.text_frame.text = "word " * 300
    footer.width, footer.height = Inches(1), Inches(0.2)

    assert fit_shape(footer) is None
    assert fit_presentation(prs) == []


def test_fit_presentation_reports_each_shrunk_shape():
    prs = Presentation()
    for _ in range(3):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _textbox(slide, "word " * 300, width_in=3, height_in=1)

    results = fit_presentation(prs)

    assert len(results) == 3
    assert [r.slide_index for r in results] == [1, 2, 3]
    assert all(r.font_scale < 1.0 for r in results)


def test_save_with_autofit_writes_scale(tmp_path):
    from pptxtpl import PptxTemplate

    template_path = tmp_path / "template.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _textbox(slide, "{{ body }}", width_in=3, height_in=1)
    prs.save(template_path)

    out_path = tmp_path / "out.pptx"
    tpl = PptxTemplate(str(template_path))
    tpl.render({"body": "word " * 300})
    tpl.save(str(out_path), autofit=True)

    rendered = Presentation(str(out_path))
    node = rendered.slides[0].shapes[0].text_frame._txBody.bodyPr.find(qn("a:normAutofit"))
    assert node is not None
    assert int(node.get("fontScale")) < 100000


def test_save_without_autofit_leaves_shape_untouched(tmp_path):
    from pptxtpl import PptxTemplate

    template_path = tmp_path / "template.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _textbox(slide, "{{ body }}", width_in=3, height_in=1)
    prs.save(template_path)

    out_path = tmp_path / "out.pptx"
    tpl = PptxTemplate(str(template_path))
    tpl.render({"body": "word " * 300})
    tpl.save(str(out_path))

    rendered = Presentation(str(out_path))
    assert rendered.slides[0].shapes[0].text_frame._txBody.bodyPr.find(qn("a:normAutofit")) is None


def test_uniform_scales_every_shape_on_the_slide_alike():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    crowded = _textbox(slide, "word " * 200, width_in=3, height_in=1)
    roomy = _textbox(slide, "word " * 20, width_in=3, height_in=1)

    fit_slide(slide, uniform=True)

    assert _norm_autofit(crowded).get("fontScale") == _norm_autofit(roomy).get("fontScale")


def test_uniform_takes_the_smallest_scale_any_shape_needs():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    crowded = _textbox(slide, "word " * 200, width_in=3, height_in=1)
    _textbox(slide, "word " * 20, width_in=3, height_in=1)

    alone = Presentation()
    alone_slide = alone.slides.add_slide(alone.slide_layouts[6])
    reference = _textbox(alone_slide, "word " * 200, width_in=3, height_in=1)
    fit_shape(reference)

    fit_slide(slide, uniform=True)

    assert _norm_autofit(crowded).get("fontScale") == _norm_autofit(reference).get("fontScale")


def test_grow_expands_a_cramped_shape_into_free_space():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = _textbox(slide, "word " * 40, width_in=6, height_in=0.4)
    original_height = box.height

    without = fit_shape(box)
    fit_slide(slide, grow=True, slide_height_emu=prs.slide_height)

    assert box.height > original_height
    # A missing fontScale means no shrink was needed at all.
    raw = _norm_autofit(box).get("fontScale")
    assert (float(raw) / 100000 if raw else 1.0) > without


def test_grow_stops_at_the_shape_below():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    top = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(6), Inches(0.4))
    top.text_frame.word_wrap = True
    top.text_frame.text = "word " * 40
    for run in top.text_frame.paragraphs[0].runs:
        run.font.size = Pt(18)
    blocker = slide.shapes.add_textbox(Inches(0.5), Inches(2.0), Inches(6), Inches(1))
    blocker.text_frame.text = "below"

    fit_slide(slide, grow=True, slide_height_emu=prs.slide_height)

    assert top.top + top.height <= blocker.top


def test_grow_leaves_a_shape_with_nothing_to_gain_alone():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = _textbox(slide, "Hi", width_in=6, height_in=2)
    original_height = box.height

    fit_slide(slide, grow=True, slide_height_emu=prs.slide_height)

    assert box.height == original_height


def test_grow_requires_a_slide_height():
    from pptxtpl.exceptions import PptxTemplateError

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _textbox(slide, "word " * 50, width_in=3, height_in=0.5)

    with pytest.raises(PptxTemplateError):
        fit_slide(slide, grow=True)


# --- review items 1 and 2 -------------------------------------------------


def test_each_word_is_measured_once_per_size(monkeypatch):
    # Re-measuring the whole growing line for every word is quadratic in the
    # words per line, and the binary search repeats it ~13 times.
    from PIL import ImageFont

    calls = []
    original = ImageFont.FreeTypeFont.getlength

    def counting(self, text, *args, **kwargs):
        calls.append(text)
        return original(self, text, *args, **kwargs)

    monkeypatch.setattr(ImageFont.FreeTypeFont, "getlength", counting)

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = _textbox(slide, "lorem ipsum dolor sit amet " * 40, width_in=8, height_in=1)

    assert fit_shape(box) < 1.0
    # 5 distinct words + a space, at no more than 13 distinct sizes.
    assert len(calls) <= 6 * 13


def test_fitting_textbox_keeps_sp_autofit():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = _textbox(slide, "Hi", width_in=8, height_in=3)
    assert _body_pr(box).find(qn("a:spAutoFit")) is not None

    assert fit_shape(box) == 1.0

    assert _body_pr(box).find(qn("a:spAutoFit")) is not None
    assert _norm_autofit(box) is None


def test_fitting_shape_with_no_autofit_is_left_alone():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = _textbox(slide, "Hi", width_in=8, height_in=3)
    body_pr = _body_pr(box)
    body_pr.remove(body_pr.find(qn("a:spAutoFit")))
    body_pr.append(body_pr.makeelement(qn("a:noAutofit"), {}))

    fit_shape(box)

    assert body_pr.find(qn("a:noAutofit")) is not None
    assert _norm_autofit(box) is None


def test_stale_font_scale_is_cleared_when_text_now_fits():
    # A template saved by PowerPoint may carry a fontScale from its old text.
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = _textbox(slide, "Hi", width_in=8, height_in=3)
    body_pr = _body_pr(box)
    body_pr.remove(body_pr.find(qn("a:spAutoFit")))
    body_pr.append(
        body_pr.makeelement(qn("a:normAutofit"), {"fontScale": "62500", "lnSpcReduction": "10000"})
    )

    fit_shape(box)

    node = _norm_autofit(box)
    assert node is not None
    assert node.get("fontScale") is None
    assert node.get("lnSpcReduction") is None


def test_uniform_leaves_fitting_shapes_autofit_alone_when_nothing_shrinks():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    a = _textbox(slide, "Hi", width_in=8, height_in=1)
    b = _textbox(slide, "There", width_in=8, height_in=1)

    assert fit_slide(slide, uniform=True) == []

    assert _body_pr(a).find(qn("a:spAutoFit")) is not None
    assert _body_pr(b).find(qn("a:spAutoFit")) is not None


# --- review items 3, 4 and minors --------------------------------------------


def test_unwrapped_frame_is_measured_as_one_line():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = _textbox(slide, "word " * 20, width_in=3, height_in=0.5)
    box.text_frame.word_wrap = False

    assert fit_shape(box) == 1.0


def test_unwrapped_frame_still_counts_hard_breaks():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = _textbox(slide, "line", width_in=8, height_in=0.5)
    box.text_frame.word_wrap = False
    paragraph = box.text_frame.paragraphs[0]
    for _ in range(10):
        paragraph.add_line_break()
        run = paragraph.add_run()
        run.text = "line"
        run.font.size = Pt(18)

    assert fit_shape(box) < 1.0


def test_snap_down_is_immune_to_float_error():
    from pptxtpl.autofit import _snap_down

    assert _snap_down(0.3) == pytest.approx(0.3)
    assert _snap_down(0.7) == pytest.approx(0.7)
    assert _snap_down(0.31) == pytest.approx(0.3)
    assert _snap_down(0.3249) == pytest.approx(0.3)


def test_autofit_types_are_exported():
    from pptxtpl import AutofitError, AutofitResult
    from pptxtpl.exceptions import PptxTemplateError

    assert issubclass(AutofitError, PptxTemplateError)
    assert AutofitResult._fields == ("slide_index", "shape_name", "font_scale")


def test_save_rejects_autofit_options_without_autofit(tmp_path):
    from pptxtpl import PptxTemplate

    template_path = tmp_path / "template.pptx"
    Presentation().save(template_path)
    tpl = PptxTemplate(str(template_path))
    tpl.render({})

    with pytest.raises(TypeError, match="autofit"):
        tpl.save(str(tmp_path / "out.pptx"), uniform=True)


# --- review items 7, 8, 9 and inset inheritance ------------------------------


def _set_lvl_size(text_frame, level, sz):
    from lxml import etree

    lst = text_frame._txBody.find(qn("a:lstStyle"))
    lvl = lst.find(qn(f"a:lvl{level + 1}pPr"))
    if lvl is None:
        lvl = etree.SubElement(lst, qn(f"a:lvl{level + 1}pPr"))
    d = lvl.find(qn("a:defRPr"))
    if d is None:
        d = etree.SubElement(lvl, qn("a:defRPr"))
    d.set("sz", str(sz))


def test_layout_placeholder_is_matched_by_idx_not_type():
    from pptxtpl.autofit import _inherited_size_pt

    prs = Presentation()
    layout = prs.slide_layouts[3]  # Two Content: idx 1 and 2 are both OBJECT
    by_idx = {p.placeholder_format.idx: p for p in layout.placeholders}
    _set_lvl_size(by_idx[1].text_frame, 0, 4000)
    _set_lvl_size(by_idx[2].text_frame, 0, 1200)
    slide = prs.slides.add_slide(layout)
    on_slide = {p.placeholder_format.idx: p for p in slide.placeholders}

    assert _inherited_size_pt(on_slide[1], 0, 18.0) == 40.0
    assert _inherited_size_pt(on_slide[2], 0, 18.0) == 12.0


def test_paragraph_level_selects_its_own_inherited_size():
    from pptxtpl.autofit import _inherited_size_pt

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    body = slide.placeholders[1]

    # Default template bodyStyle: lvl1 32pt, lvl2 28pt, lvl3 24pt.
    assert _inherited_size_pt(body, 0, 18.0) == 32.0
    assert _inherited_size_pt(body, 1, 18.0) == 28.0
    assert _inherited_size_pt(body, 2, 18.0) == 24.0


def test_indented_paragraphs_are_measured_at_their_level_size():
    def body_with(level):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        body = slide.placeholders[1]
        body.text_frame.text = "word " * 60
        for _ in range(6):
            p = body.text_frame.add_paragraph()
            p.text = "word " * 60
        for p in body.text_frame.paragraphs:
            p.level = level
        return body

    assert fit_shape(body_with(2)) > fit_shape(body_with(0))


def test_grow_skips_frame_with_inherited_middle_anchor():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    assert prs.slide_master.placeholders[0].text_frame._txBody.bodyPr.get("anchor") == "ctr"
    assert title.text_frame.vertical_anchor is None
    title.text_frame.text = "A long title that wraps onto several lines in the title box " * 3
    original_height = title.height

    fit_slide(slide, grow=True, slide_height_emu=prs.slide_height)

    assert title.height == original_height


def test_insets_are_inherited_from_the_master():
    def body_scale(master_inset_emu):
        prs = Presentation()
        master_body = prs.slide_master.placeholders[1]
        for name in ("lIns", "rIns"):
            master_body.text_frame._txBody.bodyPr.set(name, str(master_inset_emu))
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        body = slide.placeholders[1]
        body.text_frame.text = "word " * 150
        assert body.text_frame._txBody.bodyPr.get("lIns") is None
        return fit_shape(body)

    assert body_scale(Inches(3)) < body_scale(Inches(0.1))


# --- review items 5, 6 and 10 ------------------------------------------------


def _tall_text(slide, left_in, top_in, width_in, height_in):
    box = slide.shapes.add_textbox(Inches(left_in), Inches(top_in), Inches(width_in), Inches(height_in))
    box.text_frame.word_wrap = True
    box.text_frame.text = "word " * 40
    for run in box.text_frame.paragraphs[0].runs:
        run.font.size = Pt(18)
    return box


def test_grow_stops_at_layout_graphics():
    prs = Presentation()
    layout = prs.slide_layouts[6]
    # LayoutShapes has no add_shape; build the logo on a slide and move it.
    scratch = prs.slides.add_slide(layout)
    logo = scratch.shapes.add_shape(1, Inches(0.5), Inches(3.0), Inches(6), Inches(0.5))
    layout.shapes._spTree.append(logo._element)
    slide = prs.slides.add_slide(layout)
    box = _tall_text(slide, 0.5, 0.5, 6, 0.4)

    fit_slide(slide, grow=True, slide_height_emu=prs.slide_height)

    assert box.height > Inches(0.4)
    assert box.top + box.height <= Inches(3.0)


def test_grow_stops_at_master_graphics():
    prs = Presentation()
    layout = prs.slide_layouts[6]
    scratch = prs.slides.add_slide(layout)
    rule = scratch.shapes.add_shape(1, Inches(0.5), Inches(3.0), Inches(6), Inches(0.1))
    prs.slide_master.shapes._spTree.append(rule._element)
    slide = prs.slides.add_slide(layout)
    box = _tall_text(slide, 0.5, 0.5, 6, 0.4)

    fit_slide(slide, grow=True, slide_height_emu=prs.slide_height)

    assert box.height > Inches(0.4)
    assert box.top + box.height <= Inches(3.0)


def test_grow_ignores_master_graphics_hidden_by_the_layout():
    prs = Presentation()
    layout = prs.slide_layouts[6]
    layout._element.set("showMasterSp", "0")
    scratch = prs.slides.add_slide(layout)
    rule = scratch.shapes.add_shape(1, Inches(0.5), Inches(3.0), Inches(6), Inches(0.1))
    prs.slide_master.shapes._spTree.append(rule._element)
    slide = prs.slides.add_slide(layout)
    box = _tall_text(slide, 0.5, 0.5, 6, 0.4)

    fit_slide(slide, grow=True, slide_height_emu=prs.slide_height)

    assert box.top + box.height > Inches(3.0)


def test_grow_ignores_layout_placeholders_as_obstacles():
    # Layout placeholders are not rendered on the slide, so they cannot block.
    prs = Presentation()
    layout = prs.slide_layouts[1]  # body placeholder sits under the title
    slide = prs.slides.add_slide(layout)
    for ph in list(slide.placeholders):
        ph._element.getparent().remove(ph._element)
    box = _tall_text(slide, 0.5, 0.5, 6, 0.4)
    body_top = next(p for p in layout.placeholders if p.placeholder_format.idx == 1).top

    fit_slide(slide, grow=True, slide_height_emu=prs.slide_height)

    assert box.top + box.height > body_top


def test_grow_leaves_a_bottom_margin():
    from pptxtpl.autofit import GROW_BOTTOM_MARGIN_EMU

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = _tall_text(slide, 0.5, 0.5, 6, 0.4)

    fit_slide(slide, grow=True, slide_height_emu=prs.slide_height)

    assert box.top + box.height <= prs.slide_height - GROW_BOTTOM_MARGIN_EMU
    assert box.top + box.height > prs.slide_height - 2 * GROW_BOTTOM_MARGIN_EMU


def test_grow_does_not_touch_shapes_inside_groups():
    # Group children live in the group's child coordinate space, which cannot
    # be compared against slide-level obstacles.
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    child = group.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(0.4))
    child.text_frame.word_wrap = True
    child.text_frame.text = "word " * 40
    original_height = child.height

    results = fit_slide(slide, grow=True, slide_height_emu=prs.slide_height)

    assert child.height == original_height
    assert results and results[0][0] == child.name  # still shrunk


def test_grow_revert_does_not_pin_a_placeholder_xfrm():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    body = slide.placeholders[1]
    body.text_frame.text = "\n".join(["word " * 60] * 40)  # cannot fit even when grown
    assert body._element.spPr.find(qn("a:xfrm")) is None
    original_height = body.height

    fit_slide(slide, grow=True, slide_height_emu=prs.slide_height)

    assert body.height == original_height
    assert body._element.spPr.find(qn("a:xfrm")) is None


def test_word_wider_than_the_frame_wraps_mid_word():
    from pptxtpl.autofit import _wrapped_segment_count, _word_width

    word = "Supercalifragilisticexpialidocious" * 3
    width = _word_width(18, None, word)

    assert _wrapped_segment_count(word, 18, None, width / 4) >= 4
    assert _wrapped_segment_count(word, 18, None, width * 2) == 1


def test_overlong_url_in_a_short_box_is_shrunk():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = _textbox(slide, "https://example.com/" + "segment/" * 40, width_in=3, height_in=0.5)

    assert fit_shape(box) < 1.0
