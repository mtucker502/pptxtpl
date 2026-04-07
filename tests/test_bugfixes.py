"""Tests for identified bug fixes."""

import os

import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from pptxtpl import PptxTemplate
from pptxtpl.exceptions import InvalidTemplateError, TemplateRenderError
from pptxtpl.slide_ops import _remap_rids, delete_slide
from pptxtpl.xml_utils import clean_entities_in_tags


# ---------- Bug 1: _remap_rids chained-rename collision ----------
def test_remap_rids_no_chained_collision():
    xml = (
        '<root xmlns:r="http://x">'
        '<a r:id="rId2"/><b r:id="rId3"/>'
        "</root>"
    )
    el = etree.fromstring(xml)
    _remap_rids(el, {"rId2": "rId3", "rId3": "rId4"})
    children = list(el)
    assert children[0].get("{http://x}id") == "rId3"
    assert children[1].get("{http://x}id") == "rId4"


# ---------- Bug 2: delete_slide leaves orphan notes part ----------
def _add_slide_with_notes(prs, text, notes):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    tb.text_frame.text = text
    s.notes_slide.notes_text_frame.text = notes
    return s


def test_delete_slide_drops_notes_part(tmp_dir):
    prs = Presentation()
    _add_slide_with_notes(prs, "S1", "n1")
    _add_slide_with_notes(prs, "S2", "n2")
    # capture the notes part of slide 2
    slide = prs.slides[1]
    notes_part = slide.notes_slide.part
    # Before delete: notes part has the slide as a target via SLIDE rel
    delete_slide(prs, 1)
    # The notes part's back-rel to the deleted slide must be gone:
    # check that after delete the notes part is no longer reachable as a
    # rel target from any other part in the package.
    pkg_parts = list(prs.part.package.iter_parts())
    referenced = set()
    for p in pkg_parts:
        for rel in p.rels.values():
            if not rel.is_external:
                referenced.add(id(rel.target_part))
    assert id(notes_part) not in referenced
    out = os.path.join(tmp_dir, "del.pptx")
    prs.save(out)
    prs2 = Presentation(out)
    notes_parts = [
        p for p in prs2.part.package.iter_parts()
        if "notesSlide" in p.partname and "notesSlideMaster" not in p.partname
    ]
    assert len(notes_parts) == 1


# ---------- Bug 3: template slide removal orphan in render ----------
def test_slide_loop_with_notes_no_orphan(tmp_dir):
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[6])
    tb1 = s1.shapes.add_textbox(Inches(1), Inches(0.5), Inches(5), Inches(0.5))
    tb1.text_frame.text = "{%slide for x in xs %}"
    tb1b = s1.shapes.add_textbox(Inches(1), Inches(1.5), Inches(5), Inches(1))
    tb1b.text_frame.text = "Item: {{ x }}"
    tb1c = s1.shapes.add_textbox(Inches(1), Inches(2.5), Inches(5), Inches(0.5))
    tb1c.text_frame.text = "{%slide endfor %}"
    s1.notes_slide.notes_text_frame.text = "template notes {{ x }}"
    path = os.path.join(tmp_dir, "ln.pptx")
    prs.save(path)

    tpl = PptxTemplate(path)
    tpl.render({"xs": ["a", "b"]})
    # Pre-save: check that no notes part is unreferenced
    pkg = tpl._prs.part.package
    referenced = set()
    for p in pkg.iter_parts():
        for rel in p.rels.values():
            if not rel.is_external:
                referenced.add(id(rel.target_part))
    for p in pkg.iter_parts():
        if "notesSlide" in p.partname and "notesSlideMaster" not in p.partname:
            assert id(p) in referenced, f"Orphan notes part: {p.partname}"
    out = os.path.join(tmp_dir, "ln_out.pptx")
    tpl.save(out)
    prs2 = Presentation(out)  # must reopen
    # Two slides expected
    assert len(prs2.slides) == 2
    # Each remaining notesSlide must be reachable from a slide
    reachable = set()
    for s in prs2.slides:
        try:
            reachable.add(s.notes_slide.part)
        except Exception:
            pass
    notes_parts = [
        p for p in prs2.part.package.iter_parts()
        if "notesSlide" in p.partname and "notesSlideMaster" not in p.partname
    ]
    for np in notes_parts:
        assert np in reachable


# ---------- Bug 12: nested context values escaped ----------
def test_nested_context_xml_escaped(tmp_dir):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    tb.text_frame.text = "{{ user.name }}"
    path = os.path.join(tmp_dir, "nest.pptx")
    prs.save(path)

    tpl = PptxTemplate(path)
    tpl.render({"user": {"name": "Tom & Jerry"}})
    out = os.path.join(tmp_dir, "nest_out.pptx")
    tpl.save(out)
    prs2 = Presentation(out)
    texts = []
    for shape in prs2.slides[0].shapes:
        if shape.has_text_frame:
            texts.append(shape.text_frame.text)
    assert "Tom & Jerry" in texts


# ---------- Bug 13: \n line break preserves run properties ----------
def test_newline_preserves_run_properties(tmp_dir):
    from pptx.util import Pt
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2))
    p = tb.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = "{{ var }}"
    run.font.bold = True
    run.font.size = Pt(24)
    path = os.path.join(tmp_dir, "br.pptx")
    prs.save(path)

    tpl = PptxTemplate(path)
    tpl.render({"var": "line1\nline2"})
    out = os.path.join(tmp_dir, "br_out.pptx")
    tpl.save(out)
    prs2 = Presentation(out)
    runs = []
    for shape in prs2.slides[0].shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                for r in para.runs:
                    runs.append(r)
    # Find runs with the two texts
    line1 = [r for r in runs if "line1" in r.text]
    line2 = [r for r in runs if "line2" in r.text]
    assert line1 and line2
    assert line1[0].font.bold is True
    assert line2[0].font.bold is True
    # Also confirm <a:br/> present in xml
    xml = etree.tostring(prs2.slides[0]._element, encoding="unicode")
    assert "a:br" in xml


# ---------- Bug 15: get_undeclared raises on syntax error ----------
def test_get_undeclared_raises_on_syntax_error(tmp_dir):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    tb.text_frame.text = "{% if foo %}unclosed"
    path = os.path.join(tmp_dir, "synerr.pptx")
    prs.save(path)
    tpl = PptxTemplate(path)
    with pytest.raises(TemplateRenderError):
        tpl.get_undeclared_template_variables()


# ---------- Bug 8: multiple {%slide if%} on one slide ----------
def test_multiple_slide_if_raises(tmp_dir):
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s1.shapes.add_textbox(Inches(1), Inches(0.5), Inches(5), Inches(0.5))
    tb.text_frame.text = "{%slide if a %}"
    tb2 = s1.shapes.add_textbox(Inches(1), Inches(1.5), Inches(5), Inches(0.5))
    tb2.text_frame.text = "{%slide if b %}"
    tb3 = s1.shapes.add_textbox(Inches(1), Inches(2.5), Inches(5), Inches(0.5))
    tb3.text_frame.text = "{%slide endif %}"
    path = os.path.join(tmp_dir, "multif.pptx")
    prs.save(path)
    tpl = PptxTemplate(path)
    with pytest.raises(InvalidTemplateError):
        tpl.render({"a": True, "b": True})


# ---------- Bug 9: {%slide if%} inside {%slide for%} raises ----------
def test_slide_if_inside_slide_for_raises(tmp_dir):
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[6])
    s1.shapes.add_textbox(Inches(1), Inches(0.5), Inches(5), Inches(0.5)).text_frame.text = "title"
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    s2.shapes.add_textbox(Inches(1), Inches(0.5), Inches(5), Inches(0.5)).text_frame.text = "{%slide for x in xs %}"
    s3 = prs.slides.add_slide(prs.slide_layouts[6])
    s3.shapes.add_textbox(Inches(1), Inches(0.5), Inches(5), Inches(0.5)).text_frame.text = "{%slide if x %}content{%slide endif %}"
    s4 = prs.slides.add_slide(prs.slide_layouts[6])
    s4.shapes.add_textbox(Inches(1), Inches(0.5), Inches(5), Inches(0.5)).text_frame.text = "{%slide endfor %}"
    path = os.path.join(tmp_dir, "ifinfor.pptx")
    prs.save(path)
    tpl = PptxTemplate(path)
    with pytest.raises(InvalidTemplateError):
        tpl.render({"xs": [1, 2]})


# ---------- Bug 20: clean_entities_in_tags handles numeric entities ----------
def test_clean_entities_numeric():
    xml = "<a:t>{{ foo &#39;bar&#39; &#x27;baz&#x27; }}</a:t>"
    out = clean_entities_in_tags(xml)
    assert "&#39;" not in out
    assert "&#x27;" not in out
    assert "'bar'" in out
    assert "'baz'" in out
