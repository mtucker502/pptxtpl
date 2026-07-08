"""End-to-end tests for nested {%slide for%} loops."""

import os

import pytest
from pptx import Presentation
from pptx.util import Inches

from pptxtpl import PptxTemplate
from pptxtpl.exceptions import InvalidTemplateError


def _get_slide_text(slide):
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            texts.append(shape.text_frame.text)
    return " ".join(texts)


def _add_slide(prs, *texts):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    for j, text in enumerate(texts):
        tb = s.shapes.add_textbox(
            Inches(1), Inches(0.5 + j), Inches(8), Inches(0.8)
        )
        tb.text_frame.text = text
    return s


REGIONS = [
    {"name": "West", "cities": [{"name": "SF"}, {"name": "LA"}]},
    {"name": "East", "cities": [{"name": "NYC"}]},
]


class TestNestedSingleSlideInner:
    def test_inner_single_slide_loop_in_multi_slide_outer(self, tmp_dir):
        prs = Presentation()
        # Slide 1: outer for + region header
        _add_slide(
            prs,
            "{%slide for region in regions %}",
            "Region: {{ region.name }}",
        )
        # Slide 2: inner single-slide loop (both tags on this slide)
        _add_slide(
            prs,
            "{%slide for city in region.cities %}",
            "City: {{ city.name }} in {{ region.name }}",
            "{%slide endfor %}",
        )
        # Slide 3: outer endfor + region footer
        _add_slide(
            prs, "End of {{ region.name }}", "{%slide endfor %}"
        )
        path = os.path.join(tmp_dir, "nested.pptx")
        prs.save(path)

        tpl = PptxTemplate(path)
        tpl.render({"regions": REGIONS})
        output = os.path.join(tmp_dir, "out.pptx")
        tpl.save(output)

        result = Presentation(output)
        texts = [_get_slide_text(s) for s in result.slides]
        # West: header, SF, LA, footer; East: header, NYC, footer = 7
        assert len(texts) == 7
        assert "Region: West" in texts[0]
        assert "City: SF in West" in texts[1]
        assert "City: LA in West" in texts[2]
        assert "End of West" in texts[3]
        assert "Region: East" in texts[4]
        assert "City: NYC in East" in texts[5]
        assert "End of East" in texts[6]

    def test_no_jinja_tags_remain(self, tmp_dir):
        prs = Presentation()
        _add_slide(prs, "{%slide for r in regions %}", "{{ r.name }}")
        _add_slide(
            prs,
            "{%slide for c in r.cities %}{{ c.name }}{%slide endfor %}",
        )
        _add_slide(prs, "{%slide endfor %}")
        path = os.path.join(tmp_dir, "nested2.pptx")
        prs.save(path)

        tpl = PptxTemplate(path)
        tpl.render({"regions": REGIONS})
        output = os.path.join(tmp_dir, "out.pptx")
        tpl.save(output)

        for slide in Presentation(output).slides:
            text = _get_slide_text(slide)
            assert "{%" not in text
            assert "{{" not in text


class TestNestedEmptyIterables:
    def test_empty_inner_keeps_outer_slides(self, tmp_dir):
        prs = Presentation()
        _add_slide(
            prs, "{%slide for r in regions %}", "Region: {{ r.name }}"
        )
        _add_slide(
            prs,
            "{%slide for c in r.cities %}{{ c.name }}{%slide endfor %}",
        )
        _add_slide(prs, "Footer {{ r.name }}", "{%slide endfor %}")
        path = os.path.join(tmp_dir, "empty_inner.pptx")
        prs.save(path)

        tpl = PptxTemplate(path)
        tpl.render({
            "regions": [
                {"name": "Empty", "cities": []},
                {"name": "Full", "cities": [{"name": "X"}]},
            ],
        })
        output = os.path.join(tmp_dir, "out.pptx")
        tpl.save(output)

        texts = [_get_slide_text(s) for s in Presentation(output).slides]
        # Empty: header+footer (2); Full: header+city+footer (3)
        assert len(texts) == 5
        assert "Region: Empty" in texts[0]
        assert "Footer Empty" in texts[1]
        assert "Region: Full" in texts[2]
        assert "X" in texts[3]
        assert "Footer Full" in texts[4]

    def test_empty_outer_removes_everything(self, tmp_dir):
        prs = Presentation()
        _add_slide(prs, "Intro")
        _add_slide(prs, "{%slide for r in regions %}")
        _add_slide(
            prs,
            "{%slide for c in r.cities %}{{ c.name }}{%slide endfor %}",
        )
        _add_slide(prs, "{%slide endfor %}")
        _add_slide(prs, "Outro")
        path = os.path.join(tmp_dir, "empty_outer.pptx")
        prs.save(path)

        tpl = PptxTemplate(path)
        tpl.render({"regions": []})
        output = os.path.join(tmp_dir, "out.pptx")
        tpl.save(output)

        texts = [_get_slide_text(s) for s in Presentation(output).slides]
        assert len(texts) == 2
        assert "Intro" in texts[0]
        assert "Outro" in texts[1]


class TestNestedErrors:
    def test_crossing_tags_raise(self, tmp_dir):
        # endfor-without-for is how a crossing/unbalanced template surfaces
        prs = Presentation()
        _add_slide(prs, "{%slide endfor %}")
        _add_slide(prs, "{%slide for x in xs %}", "{%slide endfor %}")
        path = os.path.join(tmp_dir, "crossing.pptx")
        prs.save(path)

        tpl = PptxTemplate(path)
        with pytest.raises(InvalidTemplateError):
            tpl.render({"xs": [1]})

    def test_unclosed_for_raises(self, tmp_dir):
        prs = Presentation()
        _add_slide(prs, "{%slide for x in xs %}")
        path = os.path.join(tmp_dir, "unclosed.pptx")
        prs.save(path)

        tpl = PptxTemplate(path)
        with pytest.raises(InvalidTemplateError):
            tpl.render({"xs": [1]})

    def test_slide_if_inside_for_still_raises(self, tmp_dir):
        prs = Presentation()
        _add_slide(prs, "{%slide for x in xs %}")
        _add_slide(prs, "{%slide if x.show %}", "{{ x.name }}",
                   "{%slide endif %}")
        _add_slide(prs, "{%slide endfor %}")
        path = os.path.join(tmp_dir, "if_in_for.pptx")
        prs.save(path)

        tpl = PptxTemplate(path)
        with pytest.raises(InvalidTemplateError):
            tpl.render({"xs": [{"show": True, "name": "A"}]})


class TestNamedLoopHelpers:
    def test_named_helpers_across_nesting(self, tmp_dir):
        prs = Presentation()
        _add_slide(
            prs,
            "{%slide for region in regions as regionloop %}",
            "Region {{ regionloop.index }}/{{ regionloop.length }}",
        )
        _add_slide(
            prs,
            "{%slide for city in region.cities as cityloop %}",
            "City {{ cityloop.index }}/{{ cityloop.length }}"
            " of region {{ regionloop.index }}",
            "{%slide endfor %}",
        )
        _add_slide(prs, "{%slide endfor %}")
        path = os.path.join(tmp_dir, "named.pptx")
        prs.save(path)

        tpl = PptxTemplate(path)
        tpl.render({"regions": REGIONS})
        output = os.path.join(tmp_dir, "out.pptx")
        tpl.save(output)

        texts = [_get_slide_text(s) for s in Presentation(output).slides]
        # West: header, SF, LA, endfor-slide; East: header, NYC, endfor-slide
        assert "Region 1/2" in texts[0]
        assert "City 1/2 of region 1" in texts[1]
        assert "City 2/2 of region 1" in texts[2]
        assert "Region 2/2" in texts[4]
        assert "City 1/1 of region 2" in texts[5]

    def test_loop_refers_to_innermost(self, tmp_dir):
        prs = Presentation()
        _add_slide(prs, "{%slide for r in regions as rloop %}",
                   "outer loop.index={{ loop.index }}")
        _add_slide(
            prs,
            "{%slide for c in r.cities %}"
            "inner loop.index={{ loop.index }} rloop={{ rloop.index }}"
            "{%slide endfor %}",
        )
        _add_slide(prs, "{%slide endfor %}")
        path = os.path.join(tmp_dir, "innermost.pptx")
        prs.save(path)

        tpl = PptxTemplate(path)
        tpl.render({"regions": REGIONS})
        output = os.path.join(tmp_dir, "out.pptx")
        tpl.save(output)

        texts = [_get_slide_text(s) for s in Presentation(output).slides]
        assert "outer loop.index=1" in texts[0]
        assert "inner loop.index=2 rloop=1" in texts[2]  # LA slide

    def test_named_helper_not_shadowed_by_inline_for(self, tmp_dir):
        """Inline Jinja {% for %} shadows `loop` but not a named helper."""
        prs = Presentation()
        _add_slide(
            prs,
            "{%slide for jsa in jsas as jsaloop %}"
            "{% for p in jsa.products %}"
            "[{{ loop.index }}:{{ jsaloop.index }}:{{ p }}]"
            "{% endfor %}"
            "{%slide endfor %}",
        )
        path = os.path.join(tmp_dir, "shadow.pptx")
        prs.save(path)

        tpl = PptxTemplate(path)
        tpl.render({
            "jsas": [
                {"products": ["a", "b"]},
                {"products": ["c"]},
            ],
        })
        output = os.path.join(tmp_dir, "out.pptx")
        tpl.save(output)

        texts = [_get_slide_text(s) for s in Presentation(output).slides]
        # inline loop.index counts products; jsaloop.index counts slides
        assert "[1:1:a]" in texts[0]
        assert "[2:1:b]" in texts[0]
        assert "[1:2:c]" in texts[1]

    def test_single_level_loop_with_as(self, tmp_dir):
        """`as name` works without nesting too."""
        prs = Presentation()
        _add_slide(
            prs,
            "{%slide for x in xs as xloop %}"
            "{{ x }} ({{ xloop.index }} of {{ xloop.length }})"
            "{%slide endfor %}",
        )
        path = os.path.join(tmp_dir, "flat_as.pptx")
        prs.save(path)

        tpl = PptxTemplate(path)
        tpl.render({"xs": ["A", "B"]})
        output = os.path.join(tmp_dir, "out.pptx")
        tpl.save(output)

        texts = [_get_slide_text(s) for s in Presentation(output).slides]
        assert "A (1 of 2)" in texts[0]
        assert "B (2 of 2)" in texts[1]


class TestNamedHelperErrors:
    def test_helper_named_loop_raises(self, tmp_dir):
        prs = Presentation()
        _add_slide(
            prs,
            "{%slide for x in xs as loop %}{{ x }}{%slide endfor %}",
        )
        path = os.path.join(tmp_dir, "bad_name.pptx")
        prs.save(path)

        tpl = PptxTemplate(path)
        with pytest.raises(InvalidTemplateError):
            tpl.render({"xs": [1]})

    def test_helper_colliding_with_var_raises(self, tmp_dir):
        prs = Presentation()
        _add_slide(
            prs,
            "{%slide for x in xs as x %}{{ x }}{%slide endfor %}",
        )
        path = os.path.join(tmp_dir, "collide.pptx")
        prs.save(path)

        tpl = PptxTemplate(path)
        with pytest.raises(InvalidTemplateError):
            tpl.render({"xs": [1]})


class TestUndeclaredVariablesWithAs:
    def test_as_tag_is_stripped_for_variable_discovery(self, tmp_dir):
        """get_undeclared_template_variables must not choke on `as name`."""
        prs = Presentation()
        _add_slide(
            prs,
            "{%slide for x in xs as xloop %}"
            "{{ x }} {{ xloop.index }}"
            "{%slide endfor %}",
        )
        path = os.path.join(tmp_dir, "vars.pptx")
        prs.save(path)

        tpl = PptxTemplate(path)
        variables = tpl.get_undeclared_template_variables()
        # Discovery is approximate (tags are stripped), but it must not raise
        assert isinstance(variables, set)
